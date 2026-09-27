import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def value(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def row_values(headers, row):
    return {header: value(cell) for header, cell in zip(headers, row)}


def export_workbook(path: Path, output: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    with path.open("rb") as source:
        source_hash = hashlib.sha256(source.read()).hexdigest()
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            sheets[sheet_name] = []
            continue
        first = rows[0]
        if sheet_name == "Dashboard" or not any(first):
            headers = [f"column_{index + 1}" for index in range(len(first))]
            data_rows = rows
            start_row = 1
        else:
            headers = [str(cell).strip() if cell not in (None, "") else f"column_{index + 1}" for index, cell in enumerate(first)]
            data_rows = rows[1:]
            start_row = 2
        sheets[sheet_name] = [
            {"rowId": index, "values": row_values(headers, row)}
            for index, row in enumerate(data_rows, start=start_row)
            if any(cell not in (None, "") for cell in row)
        ]
    output.write_text(
        json.dumps(
            {
                "filename": path.name,
                "sourcePath": str(path),
                "sourceHash": source_hash,
                "sheets": sheets,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


if __name__ == "__main__":
    export_workbook(Path(sys.argv[1]), Path(sys.argv[2]))
